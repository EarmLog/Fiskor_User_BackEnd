# ============================================================
# Exportacion de datos a Excel y documentos a Word.
#   - Excel con pandas + openpyxl (planillas, reportes, libros).
#   - Word con python-docx (cotizaciones, facturas, comprobantes).
# ============================================================
import logging
from datetime import datetime
from pathlib import Path

from app.config import Configuracion

registro = logging.getLogger("fiskor")


def _asegurar_carpeta(ruta: Path) -> Path:
    ruta.mkdir(parents=True, exist_ok=True)
    return ruta


def exportar_excel(filas: list[dict], titulo: str, nombre_base: str | None = None) -> str:
    """Escribe una lista de diccionarios en un archivo .xlsx.

    Devuelve la ruta del archivo generado.
    """
    import pandas as pd

    carpeta = _asegurar_carpeta(Path(Configuracion.RUTA_EXPORTADOS) / "excel")
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = nombre_base or f"{titulo.lower().replace(' ', '_')}_{marca}"
    ruta = carpeta / f"{nombre}.xlsx"

    if not filas:
        filas = [{"Aviso": "Sin registros para exportar"}]
    dataframe = pd.DataFrame(filas)
    dataframe.to_excel(ruta, index=False, sheet_name=titulo[:28] or "Reporte")
    registro.info("Excel generado: %s", ruta)
    return str(ruta)


def generar_documento_word(titulo: str, empresa: dict, lineas: list[str], nombre_base: str | None = None) -> str:
    """Crea un documento .docx con encabezado de la empresa.

    Se usa para cotizaciones, facturas y comprobantes simples.
    """
    from docx import Document
    from docx.shared import Pt

    carpeta = _asegurar_carpeta(Path(Configuracion.RUTA_EXPORTADOS) / "word")
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre = nombre_base or f"{titulo.lower().replace(' ', '_')}_{marca}"
    ruta = carpeta / f"{nombre}.docx"

    documento = Document()
    encabezado = documento.add_heading(empresa.get("nombre", "Fiskor ERP"), level=0)
    encabezado.alignment = 1  # centrado

    subtitulo = documento.add_paragraph()
    subtitulo.alignment = 1
    corrida = subtitulo.add_run(
        f"RIF: {empresa.get('rif', '')}    |    {empresa.get('direccion', '')}"
    )
    corrida.font.size = Pt(9)

    documento.add_heading(titulo, level=2)

    for linea in lineas:
        parrafo = documento.add_paragraph(str(linea))
        parrafo.paragraph_format.space_after = Pt(2)

    documento.save(ruta)
    registro.info("Documento Word generado: %s", ruta)
    return str(ruta)


def generar_factura_word(empresa: dict, venta: dict, detalles: list[dict]) -> str:
    """Genera una factura/nota de entrega en Word."""
    from docx import Document
    from docx.shared import Pt

    carpeta = _asegurar_carpeta(Path(Configuracion.RUTA_EXPORTADOS) / "word")
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    ruta = carpeta / f"factura_{venta.get('numero_factura', marca)}.docx"

    documento = Document()
    titulo = documento.add_heading(empresa.get("nombre", "Fiskor ERP"), level=0)
    titulo.alignment = 1

    info = documento.add_paragraph()
    info.alignment = 1
    info.add_run(
        f"RIF: {empresa.get('rif', '')}\n{empresa.get('direccion', '')}\n"
        f"Telefono: {empresa.get('telefono', '')}"
    ).font.size = Pt(9)

    tipo_doc = "FACTURA" if venta.get("tipo") == "factura" else "NOTA DE ENTREGA"
    documento.add_heading(f"{tipo_doc} N° {venta.get('numero_factura', '')}", level=2)

    documento.add_paragraph(f"Cliente: {venta.get('cliente_nombre', '')}")
    documento.add_paragraph(f"RIF/Cedula: {venta.get('cliente_rif', '')}")
    documento.add_paragraph(f"Fecha: {venta.get('fecha', '')}")
    documento.add_paragraph(f"Control: {venta.get('numero_control', '')}")

    tabla = documento.add_table(rows=1, cols=4)
    tabla.style = "Light Grid Accent 1"
    encabezados = tabla.rows[0].cells
    for indice, texto in enumerate(["Descripcion", "Cant.", "Precio $", "Total $"]):
        encabezados[indice].text = texto

    for detalle in detalles:
        celdas = tabla.add_row().cells
        celdas[0].text = str(detalle.get("descripcion", ""))
        celdas[1].text = str(detalle.get("cantidad", ""))
        celdas[2].text = f"{float(detalle.get('precio_unitario', 0)):.2f}"
        celdas[3].text = f"{float(detalle.get('total', 0)):.2f}"

    documento.add_paragraph()
    resumen = documento.add_paragraph()
    resumen.add_run(f"Subtotal: ${venta.get('subtotal', 0):.2f}\n")
    resumen.add_run(f"IVA: ${venta.get('iva', 0):.2f}\n")
    resumen.add_run(f"IGTF: Bs. {venta.get('igtf', 0):.2f}\n")
    resumen.add_run(f"TOTAL Bs.: {venta.get('total_bolivares', 0):.2f}\n")
    resumen.add_run(f"TOTAL $: {venta.get('total_divisas', 0):.2f}\n")
    resumen.add_run(f"Tasa aplicada: {venta.get('tasa_aplicada', 0):.2f} Bs/USD")

    documento.save(ruta)
    registro.info("Factura Word generada: %s", ruta)
    return str(ruta)
